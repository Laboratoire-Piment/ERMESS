# -*- coding: utf-8 -*-
"""
Created on Fri Sep 22 16:18:47 2023

@author: jlegalla
"""

import pandas as pd
import numpy as np
import openpyxl         
import copy  
import timezonefinder
import ERMESS_PV_model as EPV
import ERMESS_Windmodel as EWi
import ERMESS_meteo as Eme

def find_constraint_levels(constraint_num,Constraint_level,D_Non_movable_load,D_Movable_load,storage_characteristics,prod_C,prods_U,Bounds_prod):
    max_production = prod_C/1000+np.inner(prods_U.T,Bounds_prod)/1000 
    total_load=sum(D_Non_movable_load)+sum(D_Movable_load)       
    extra_prod = np.where(max_production>D_Non_movable_load,max_production-D_Non_movable_load,0) 
    sum_extra_prod = sum(extra_prod)-sum(D_Movable_load)        
    D_Instantaneous_energy=np.minimum(max_production,D_Non_movable_load)
    D_reportable_energy=max(storage_characteristics[4,:])*sum_extra_prod
    if constraint_num==1:
        D_max_self_sufficiency = min(1,(sum(D_Instantaneous_energy)+D_reportable_energy)/total_load)
        D_max_constraint_level = D_max_self_sufficiency
    else:
        
        ## The max. self-consumption is always 1 because we can always lose artificially energy using storage
        D_max_self_consumption = 1
        D_max_constraint_level = D_max_self_consumption
    return(tuple((constraint_num ,Constraint_level,D_max_constraint_level)))

def compute_grid_prices(datetime_data,grid_price):
    grid_price=grid_price.replace(np.nan, '', regex=True)
    prices = [pd.DataFrame(data={'Datetime':datetime_data,'Hour type':'NONE','Price':None}) for i in range(len(grid_price))]
    prices_hour_type = []
    prices_num = []


    fixed_premium = []
    Overrun = []
    Selling_price = []
    
    for i in range(len(grid_price)):
        prices[i].loc[(prices[i]['Datetime'].apply(lambda x: x.dayofweek).isin([5,6])) & (prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Weekend peak hours'].split(' '))),'Hour type']='Peak'
        prices[i].loc[(prices[i]['Datetime'].apply(lambda x: x.dayofweek).isin([5,6])) & (prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Weekend full hours'].split(' '))),'Hour type']='WE full'
        prices[i].loc[(prices[i]['Datetime'].apply(lambda x: x.dayofweek).isin([5,6])) & (prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Weekend off-peak hours'].split(' '))),'Hour type']='WE off'
        prices[i].loc[(prices[i]['Datetime'].apply(lambda x: x.dayofweek).isin([0,1,2,3,4])) & (prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Workday peak hours'].split(' '))),'Hour type']='Peak'
        prices[i].loc[(prices[i]['Datetime'].apply(lambda x: x.dayofweek).isin([0,1,2,3,4])) & (prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Workday full hours'].split(' '))),'Hour type']='W full'
        prices[i].loc[(prices[i]['Datetime'].apply(lambda x: x.dayofweek).isin([0,1,2,3,4])) & (prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Workday off-peak hours'].split(' '))),'Hour type']='W off'
    
        prices[i].loc[prices[i]['Hour type']=='Peak' ,'Price']=grid_price.iloc[i]['Peak (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='WE full') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Summer months'].split(' '))) ,'Price']=grid_price.iloc[i]['Summer full hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='W full') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Summer months'].split(' '))) ,'Price']=grid_price.iloc[i]['Summer full hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='WE off') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Summer months'].split(' '))) ,'Price']=grid_price.iloc[i]['Summer off-peak hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='W off') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Summer months'].split(' '))) ,'Price']=grid_price.iloc[i]['Summer off-peak hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='WE full') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Winter months'].split(' '))) ,'Price']=grid_price.iloc[i]['Winter full hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='W full') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Winter months'].split(' '))) ,'Price']=grid_price.iloc[i]['Winter full hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='WE off') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Winter months'].split(' '))) ,'Price']=grid_price.iloc[i]['Winter off-peak hours (c€/kWh)']
        prices[i].loc[(prices[i]['Hour type']=='W off') & (prices[i]['Datetime'].apply(lambda x: str(x.month)).isin(grid_price.iloc[i]['Winter months'].split(' '))) ,'Price']=grid_price.iloc[i]['Winter off-peak hours (c€/kWh)']

        prices[i]['Price']=prices[i]['Price']*(1+grid_price.iloc[i]['TVA load'])
        prices[i]['Price']+=grid_price.iloc[i]['CSPE (c€/kWh)']
        prices[i]['Price']=prices[i]['Price']/100
        prices_hour_type.append(prices[i]['Hour type'])
        prices_num.append(prices[i]['Price'])

        fixed_premium.append( grid_price.iloc[i]['Fixed premium (€/kW)'])
        Overrun.append( grid_price.iloc[i]['Power overrun (€/kW)'])
        Selling_price.append(np.repeat(grid_price.iloc[i]['Selling base price (c€/kWh)'],len(datetime_data)))
        Selling_price[i][(prices[i]['Datetime'].apply(lambda x: str(x.hour)).isin(grid_price.iloc[i]['Selling peak hours'].split(' ')))]=grid_price.iloc[i]['Selling peak price (c€/kWh)']
        Selling_price[i]=Selling_price[i]/100
        
        #Converting into Numpy
    fixed_premium=np.array(fixed_premium,dtype=np.float64)
    Overrun=np.array(Overrun,dtype=np.float64)
    Selling_price=np.array(Selling_price,dtype=np.float64)
    prices_num=np.float64(np.vstack(prices_num))
    prices_hour_type=np.vstack(prices_hour_type).astype('U')

    return([prices_hour_type,prices_num,fixed_premium,Overrun,Selling_price])

def timeseries_interpolation(datetime_model,series_datetime,series_yvalue):
    y_values = np.float64(np.interp(np.array(datetime_model,dtype='float64'),np.array(series_datetime,dtype='float64'),np.array(series_yvalue,dtype='float64')))
    return(y_values)

def read_data(Data) :
    
    EnR_site = {"longitude":Data['Environment']['Longitude (°)'][0],"latitude":Data['Environment']['Latitude (°)'][0],"altitude":Data['Environment']['Altitude (m)'][0]}
    time_resolution = np.float64( Data['Environment']['time resolution (steps/h)'][0])

    # Meteo data
    datetime_data=Data['TimeSeries']['Datetime']
    datetime_data=pd.to_datetime(datetime_data, format="%d/%m/%Y %H:%M").round('1s')
    
    timezone_str = timezonefinder.TimezoneFinder().certain_timezone_at(lat=EnR_site['latitude'], lng=EnR_site['longitude'])
    EnR_site ["tz"]=timezone_str

    datetime_model = pd.date_range(datetime_data[0],datetime_data[len(datetime_data)-1],freq=str(int(60/time_resolution))+'min',tz=timezone_str)#.tz_localize(timezone_str)
    
    Meteo_computation = Data['Environment']['Meteo'][0]
    
    
    Production_computation = Data['Environment']['Production'][0]
    
    if (Production_computation=='automatic'):
        
        if (Meteo_computation=='manual'):
        
            PV_meteo = Data['Meteo'].loc[:, ['GHI (W/m²)','DNI (W/m²)','DHI (W/m²)','Air temperature (°C)','Wind speed (m/s)']].rename(columns={'GHI (W/m²)':'ghi','DNI (W/m²)':'dni','DHI (W/m²)':'dhi','Air temperature (°C)':'temp_air','Wind speed (m/s)':'wind_speed'}).set_index(pd.to_datetime(Data['Meteo'].loc[:,'Datetime']))
            Wind_heights = Data['Meteo'].loc[np.array((5,3,6)),'Measurement height (m)'].to_numpy()
            Wind_meteo = pd.DataFrame(data=Data['Meteo'].loc[:, ['Wind speed (m/s)','Air temperature (°C)','Pressure (hPa)']].values,index=Data['Meteo']['Datetime'],columns=[np.array(['wind_speed','temperature','pressure']),Wind_heights])
            Wind_meteo['temperature']=Wind_meteo['temperature']+273.15
            Wind_meteo['pressure']=Wind_meteo['pressure']*100 #in Pa
            Datetime_prods = pd.DatetimeIndex(pd.to_datetime(Data['Meteo']['Datetime'],format='%d/%m/%Y %H:%M')).round('1s').tz_localize(timezone_str,nonexistent="shift_forward" ,ambiguous=False)
            
        elif (Meteo_computation == 'automatic' ):
            PV_meteo,Wind_meteo = Eme.import_meteo(EnR_site['latitude'],EnR_site['longitude'],EnR_site['altitude'],datetime_model[0],datetime_model[len(datetime_model)-1],timezone_str)
            Datetime_prods =PV_meteo.index

        Terrain_type =  Data['Environment']['Terrain type'][0]
        Roughness_length = 0 if Terrain_type=='off-shore' else (0.005 if Terrain_type=='no vegetation, no obstacles' else (0.005 if Terrain_type=='off-shore' else (0.03 if Terrain_type=='flat terrain, grass, isloated obstacles' else (0.1 if Terrain_type=='low crops, occasional obstacles' else (0.25 if Terrain_type=='high crops, scattered obstacles' else (0.5 if Terrain_type=='parkland, bushes, numerous obstacles' else (1 if Terrain_type=='regular large obstacles (suburbs, forest)' else (2.5 if Terrain_type=='city centre' else 0.5))))))))
        Wind_meteo[('roughness_length', 0)] = Roughness_length
        
        PV_prod = Data['PV_production_specs'].loc[:, ['Tilt (°)','Azimuth (°)','Module type','Mounting','Inverter','Module','Modules per string','Strings','Surface type']]
        WT_prod = Data['WT_production_specs'].loc[:, ['Model','Hub height (m)']].rename(columns={'Hub height (m)':'hub_height','Model':'turbine_type'})
        WT_prod = [EWi.default_wind_turbines(WT_prod.iloc[i]) for i in range(len(WT_prod))]
        
        PV_TempParam_default = {'a': -3.47,'b': -0.0594,'deltaT': 3} #SAPM coeffs
        PV_FixModMount = [{"type":"Fixed","surface_tilt":PV_prod['Tilt (°)'][i],"surface_azimuth":PV_prod['Azimuth (°)'][i],"racking_model":PV_prod['Mounting'][i]} for i in range(len(PV_prod))]
        PV_arrayParam = [{"surface_type":PV_prod['Surface type'][i],"module_type":PV_prod['Module type'][i],"modules_per_string":PV_prod['Modules per string'][i],"strings":PV_prod['Strings'][i]} for i in range(len(PV_prod))]
        LossesParam_default = {"soiling": 2,"shading": 3,"snow": 0,"mismatch": 2,"wiring": 2,"connections": 0.5,"lid": 1.5,"nameplate_rating": 1,"age": 0,"availability": 3}
        
        prods_U_PV = np.array([EPV.pvmodel(EnR_site, PV_meteo, PV_TempParam_default, PV_prod['Module'][i], PV_prod['Inverter'][i], PV_FixModMount[i], PV_arrayParam[i],LossesParam_default,False,False,False,False)[0] for i in range(len(PV_prod))])
        prods_U_PV = np.where(prods_U_PV<0,0,prods_U_PV)
        prods_U_WT = np.array([EWi.windmodel(weather=Wind_meteo,MyTurbineDict=WT_prod[i],ModelChainDict=None,csv=False,plot=False,pow_TS=False,pow_curv=False) for i in range(len(WT_prod))])
        prods_U = np.row_stack([x for x in (prods_U_PV, prods_U_WT) if x.size] or [np.array([])])
    elif (Production_computation=='manual'):   
        prods_U = Data['Unit_productions'].drop('Datetime',axis=1)
        prods_U=np.array(prods_U).T
        Datetime_prods = Data['Unit_productions']['Datetime']
        
    prods_U=np.array([timeseries_interpolation(datetime_model,  Datetime_prods, prods_U[i]) for i in range(len(prods_U))])
    Volums_prod = np.sum(prods_U,axis=1)
    
    
    PV_specs = Data['PV_production_specs'].loc[:,['Capital unit cost (€)','Operational unit cost (€/yrs)','Lifetime (years)','Capacity','eqCO2 Emissions (gCO2/kWh)','EROI','Surface group']].to_numpy(dtype=np.float64)
    WT_specs = Data['WT_production_specs'].loc[:,['Capital unit cost (€)','Operational unit cost (€/yrs)','Lifetime (years)','Capacity','eqCO2 Emissions (gCO2/kWh)','EROI','Surface group']].to_numpy(dtype=np.float64)    
    specs_num = np.row_stack((PV_specs,WT_specs))
    
    groups = [np.where(specs_num[:,6]==i)[0] for i in np.unique(specs_num[:,6])]  
    
    specs_Id=np.concatenate((Data['PV_production_specs']['Id'].to_numpy(dtype='U'),Data['WT_production_specs']['Id'].to_numpy(dtype='U')))
    specs_names=np.array(['Capital unit cost','Operational unit cost','Lifetime','Capacity','eqCO2 Emissions','EROI','Surface group'],dtype='U')
    
    series_datetime = pd.to_datetime(Data['TimeSeries']['Datetime'], format="%d/%m/%Y %H:%M").dt.tz_localize(timezone_str,nonexistent="shift_forward" ,ambiguous=False)
    prod_C = timeseries_interpolation(datetime_model, series_datetime, Data['TimeSeries']['Current_production (W)'])
    prod_C[prod_C<0]=0
    
    Non_movable_load = timeseries_interpolation(datetime_model, series_datetime, Data['TimeSeries']['Non-controllable load (kW)'])
    Y_movable_load = timeseries_interpolation(datetime_model, series_datetime, Data['TimeSeries']['Yearly movable load (kW)'])
    D_movable_load = timeseries_interpolation(datetime_model, series_datetime, Data['TimeSeries']['Daily movable load (kW)'])
      
    grid_price = Data['Grid_prices']
    Contract_Id = grid_price['Contract_Id']

    #time_resolution = (datetime_data.iloc[-1].round(freq='min')-datetime_data.iloc[1].round(freq='min'))/(len(datetime_data)-2)
    #time_resolution = np.float64(3600/time_resolution.seconds)
    Constraint = Data['Environment']['Constraint'][0]
    Connexion = Data['Environment']['Connexion'][0]
        
    if Constraint== 'Self-sufficiency':
            constraint_num = np.int64(1) 
    elif Constraint== 'Self-consumption':
            constraint_num = np.int64(2)
    elif Constraint== 'EnR fraction':
            constraint_num = np.int64(3)
    else : 
            print('No proper constraint found !')
            constraint_num = np.int64(0)
                
    Constraint_level = np.float64(Data['Environment']['Constraint level'][0])
    criterion = Data['Environment']['Optimisation criterion'][0]
       
    if criterion== 'LCOE':
            criterion_num = np.int64(1)
    elif criterion== 'Annual net benefits':
            criterion_num = np.int64(2)
    elif criterion== 'NPV':
            criterion_num = np.int64(3)
    elif criterion== 'Self-sufficiency':
            criterion_num = np.int64(4)
    elif criterion== 'Self-consumption':
            criterion_num = np.int64(5)
    elif criterion== 'Autonomy':
            criterion_num = np.int64(6)
    elif criterion== 'eqCO2 Emissions':
            criterion_num = np.int64(7)
    elif criterion== 'fossil fuel consumption':
            criterion_num = np.int64(8)
    elif criterion== 'EROI':
            criterion_num = np.int64(9)
    elif criterion== 'Energy losses':
            criterion_num = np.int64(10)
    elif criterion== 'Maximum power from grid':
            criterion_num = np.int64(11)   
    else : 
            criterion_num = np.int64(0)
            
    type_optim = Data['Environment']['type'][0]
    
    storages = Data['Storages']
    Cost_power = storages['PCS cost (€/kW)']+storages['BOP cost (€/kW)']
    storage_characteristics = storages.T[1:].to_numpy(dtype=np.float64)
    storage_techs = storages['Technology'].to_numpy(dtype='U')
    storage_characteristics = np.vstack((storage_characteristics,Cost_power),dtype=np.float64)
    storage_characteristics_names = np.hstack((storages.columns[1:].to_numpy(),'Cost power'))

    Bounds_prod = np.int64(specs_num[:,np.where(specs_names=='Capacity')[0][0]])
    n_UP=len(Bounds_prod)
    duration_years = np.float64(len(prod_C)/time_resolution/8760)
    
    DG_fuel_cost = Data['Diesel generator']['Value'][0]
    DG_lifetime = Data['Diesel generator']['Value'][1]
    DG_unit_cost = Data['Diesel generator']['Value'][2]
    DG_maintenance_cost = Data['Diesel generator']['Value'][3]
    fuel_CO2eq_emissions = Data['Diesel generator']['Value'][4]
    DG_EROI = Data['Diesel generator']['Value'][5]
    DG_fuel_consumption = np.array(Data['Diesel generator']['DG fuel consumption'][1:11],dtype=np.float64)
    
    
    if (type_optim=='pro'):
        Defined_items = np.array(('Discharge order','D_DSM_levels','Y_DSM_levels','Diesel generator'))[np.where(Data['Dispatching']['User-Defined']=='Yes')[0]]
        Discharge_order = (np.array(Data['Dispatching']['Discharge order'],dtype=np.int64)[~pd.isnull(np.array(Data['Dispatching']['Storages']))])-1
        Taking_over = np.array(Data['Dispatching']['Taking over level (%)'][0:9],dtype='float64')
        Taking_over_ext = np.array(Data['Dispatching']['DG/grid taking over level (%)'][0:9],dtype='float64')
        Taking_overs = np.row_stack((Taking_over,Taking_over_ext))
        energy_use_repartition_DSM = Data['Dispatching']['Repartition coefficients'][0]/(Data['Dispatching']['Repartition coefficients'][0]+Data['Dispatching']['Repartition coefficients'][1])
        D_DSM_minimum_levels = np.array(Data['Dispatching']['D_DSM minimum levels'][1:24],dtype='float64')
        Y_DSM_minimum_levels = np.array(Data['Dispatching']['Y_DSM minimum levels'][1:12],dtype='float64')
        DG_strategy = Data['Dispatching']['Diesel generator'][0]
        DG_min_runtime = Data['Dispatching']['Diesel generator'][1]
        DG_min_production = Data['Dispatching']['Diesel generator'][2]
    else : 
        Defined_items = ()
        Discharge_order, Taking_overs,energy_use_repartition_DSM, D_DSM_minimum_levels, Y_DSM_minimum_levels, DG_strategy, DG_min_runtime,DG_min_production = np.repeat(np.nan,8)
        
    Dispatching = [Defined_items, Discharge_order, Taking_overs, D_DSM_minimum_levels, Y_DSM_minimum_levels, DG_strategy, DG_min_runtime, DG_min_production,energy_use_repartition_DSM]
    
    (Grid_Fossil_fuel_ratio,Main_grid_emissions,Main_grid_ratio) = (np.float64(Data['Environment']['Main grid fossil fuel ratio'][0]),np.float64(Data['Environment']['Main grid emissions (gCO2/kWh)'][0]),np.float64(Data['Environment']['Main grid ratio primary over final energy'][0]))
   
    
    r_cross_init=np.float64(Data['Hyperparameters']['Initialisation values'][0])
    n_iter_init=np.int64(Data['Hyperparameters']['Initialisation values'][1])
    r_cross=np.float64(Data['Hyperparameters']['Evolution values'][0])
    n_iter,nb_ere,n_pop,n_nodes,n_core =np.int64(Data['Hyperparameters']['Evolution values'][1:6])
    Cost_sequence = np.float64(Data['Hyperparameters']['Constraint costs'])
    operator_contract = np.float64(Data['Hyperparameters']['Contract'])
    operator_production = np.float64(Data['Hyperparameters']['Production'])
    operator_storage_volumes = np.float64(Data['Hyperparameters']['Storage volume'])
    operator_storage_use = np.float64(Data['Hyperparameters']['Storage use'])
    operator_storage_power = np.float64(Data['Hyperparameters']['Storage power'])
    operator_scheduling_consistency = np.float64(Data['Hyperparameters']['Scheduling consistency'])
    operator_storage_timeserie = np.float64(Data['Hyperparameters']['Storage timeserie'])
    operator_storage_opposite = np.float64(Data['Hyperparameters']['Storage opposite moves'])
    operator_storage_transfers = np.float64(Data['Hyperparameters']['Storage transfers'])
    operator_storage_specification = np.float64(Data['Hyperparameters']['Storage specification'])
    operator_curve_smoothing = np.float64(Data['Hyperparameters']['Curve smoothing'])
    operator_constraint = np.float64(Data['Hyperparameters']['Constraint'])
    operator_long_term_consistency = np.float64(Data['Hyperparameters']['Long-term consistency'])
    operator_YDSM = np.float64(Data['Hyperparameters']['Yearly demand-side management'])
    operator_DDSM = np.float64(Data['Hyperparameters']['Daily demand-side management'])
    hyperparameters_main = {'Initialisation':(n_core,n_nodes,r_cross_init,n_pop,n_iter_init),'Evolution':(n_core,n_nodes,r_cross,n_pop,n_iter,nb_ere)}
    hyperparameters_operators_names = np.array(('Contract','Production','Storage volume','Storage use','Storage power','Scheduling consistency','Storage timeserie','Storage opposite moves','Storage transfers','Storage specification','Curve smoothing','Constraint','Long-term consistency','Yearly demand-side management','Daily demand-side management'))
    hyperparameters_operators_num = np.array((operator_contract,operator_production,operator_storage_volumes,operator_storage_use,operator_storage_power,operator_scheduling_consistency,operator_storage_timeserie,operator_storage_opposite,operator_storage_transfers,operator_storage_specification,operator_curve_smoothing,operator_constraint,operator_long_term_consistency,operator_YDSM,operator_DDSM)).T

    r_cross_pro = np.float64(Data['Hyperparameters_pro']['Evolution values'][0])
    n_iter_pro,n_pop_pro = np.int64(Data['Hyperparameters_pro']['Evolution values'][1:3])
    operator_contract_pro = np.float64(Data['Hyperparameters_pro']['Contract'])
    operator_production_pro = np.float64(Data['Hyperparameters_pro']['Production'])
    operator_strategy_pro = np.float64(Data['Hyperparameters_pro']['Strategy'])
    operator_discharge_order_pro = np.float64(Data['Hyperparameters_pro']['Discharge order'])
    operator_energy_use_pro = np.float64(Data['Hyperparameters_pro']['Energy use'])
    operator_taking_over_pro = np.float64(Data['Hyperparameters_pro']['Taking over'])
    operator_DSM_min_levels_pro = np.float64(Data['Hyperparameters_pro']['DSM minimum levels'])
    operator_DG_min_runtime_pro = np.float64(Data['Hyperparameters_pro']['DG min runtime'])
    operator_DG_min_production_pro = np.float64(Data['Hyperparameters_pro']['DG min production'])
    operator_storages_capacity_pro = np.float64(Data['Hyperparameters_pro']['storages capacity'])
    operator_storages_power_pro = np.float64(Data['Hyperparameters_pro']['storages power'])
    operator_init_SOC_pro = np.float64(Data['Hyperparameters_pro']['Initial SOC'])
    
    hyperparameters_main_pro = (r_cross_pro,n_pop_pro,n_iter_pro)
    hyperparameters_operators_names_pro = np.array(('Contract','Production','Strategy','Discharge order','Energy use','Taking over','DSM minimum levels','DG min runtime','DG min production','storages capacity','storages power','Initial SOC'))
    hyperparameters_operators_num_pro = np.array((operator_contract_pro,operator_production_pro,operator_strategy_pro,operator_discharge_order_pro,operator_energy_use_pro,operator_taking_over_pro,operator_DSM_min_levels_pro,operator_DG_min_runtime_pro,operator_DG_min_production_pro,operator_storages_capacity_pro,operator_storages_power_pro,operator_init_SOC_pro)).T

    [prices_hour_type,prices_num,fixed_premium,Overrun,Selling_price] = compute_grid_prices(datetime_model,grid_price)
    return (pd.Series(datetime_model),Grid_Fossil_fuel_ratio,Main_grid_emissions,Main_grid_ratio,Connexion,specs_names,specs_Id,specs_num,prod_C,prods_U,Volums_prod,Y_movable_load,D_movable_load,Non_movable_load,time_resolution,constraint_num, Constraint_level ,criterion_num,storage_techs,storage_characteristics_names,storage_characteristics,Bounds_prod,n_UP,duration_years,prices_hour_type,prices_num,fixed_premium,Overrun,Selling_price,Contract_Id,hyperparameters_main,hyperparameters_operators_names,hyperparameters_operators_num,hyperparameters_main_pro,hyperparameters_operators_names_pro,hyperparameters_operators_num_pro,Cost_sequence,type_optim,Dispatching,DG_fuel_cost,DG_lifetime,DG_unit_cost,DG_maintenance_cost,DG_fuel_consumption,DG_EROI,fuel_CO2eq_emissions,groups)    



def NONJIT_adaptation_hyperparameters_initialisation(hyperparameters_operators):
    adaptated_hyperparameters_operators=copy.deepcopy(hyperparameters_operators)
    adaptated_hyperparameters_operators['Contract'][0]=1*adaptated_hyperparameters_operators['Contract'][0]
    adaptated_hyperparameters_operators['Production'][0]=1*adaptated_hyperparameters_operators['Production'][0]
    adaptated_hyperparameters_operators['Storage volume'][0]=0*adaptated_hyperparameters_operators['Storage volume'][0]
    adaptated_hyperparameters_operators['Storage power'][0]=0*adaptated_hyperparameters_operators['Storage power'][0]
    adaptated_hyperparameters_operators['Storage use'][0]=1*adaptated_hyperparameters_operators['Storage use'][0]
    adaptated_hyperparameters_operators['Storage timeserie'][0]=1*adaptated_hyperparameters_operators['Storage timeserie'][0]
    adaptated_hyperparameters_operators['Storage opposite'][0]=1*adaptated_hyperparameters_operators['Storage opposite'][0]
    adaptated_hyperparameters_operators['Storage transfer'][0]=1*adaptated_hyperparameters_operators['Storage transfer'][0]
    adaptated_hyperparameters_operators['Long-term consistency'][0]=0*adaptated_hyperparameters_operators['Long-term consistency'][0]
    adaptated_hyperparameters_operators['Scheduling consistency'][0]=0.2*adaptated_hyperparameters_operators['Scheduling consistency'][0]
    adaptated_hyperparameters_operators['Storage specification'][0]=0*adaptated_hyperparameters_operators['Storage specification'][0]
    adaptated_hyperparameters_operators['Curve smoothing'][0]=0.7*adaptated_hyperparameters_operators['Curve smoothing'][0]
    adaptated_hyperparameters_operators['Y_DSM'][0]=adaptated_hyperparameters_operators['Y_DSM'][0]
    adaptated_hyperparameters_operators['D_DSM'][0]=adaptated_hyperparameters_operators['D_DSM'][0]

    return (adaptated_hyperparameters_operators)

def adaptation_hyperparameters_initialisation(hyperparameters_operators_num,hyperparameters_operators_names):
    adaptated_hyperparameters_operators=copy.deepcopy(hyperparameters_operators_num)
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Contract']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Contract'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Production']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Production'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage volume']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage volume'][0][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage power']=0.2*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage power'][0]
    #adaptated_hyperparameters_operators[1,hyperparameters_operators_names=='Storage power']=40
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage use']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage use'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage timeserie']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage timeserie'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage opposite']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage opposite'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage transfers']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage transfers'][0]
    adaptated_hyperparameters_operators[1,hyperparameters_operators_names=='Storage transfers']=5
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Long-term consistency']=0*hyperparameters_operators_num[:,hyperparameters_operators_names=='Long-term consistency'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Scheduling consistency']=0.5*hyperparameters_operators_num[:,hyperparameters_operators_names=='Scheduling consistency'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Storage specification']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Storage specification'][0]
    adaptated_hyperparameters_operators[1,hyperparameters_operators_names=='Storage specification']=1
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Constraint']=1*hyperparameters_operators_num[:,hyperparameters_operators_names=='Constraint'][0]
    adaptated_hyperparameters_operators[1,hyperparameters_operators_names=='Constraint']=20
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Curve smoothing']=0.7*hyperparameters_operators_num[:,hyperparameters_operators_names=='Curve smoothing'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='Y_DSM']=hyperparameters_operators_num[:,hyperparameters_operators_names=='Y_DSM'][0]
    adaptated_hyperparameters_operators[0,hyperparameters_operators_names=='D_DSM']=hyperparameters_operators_num[:,hyperparameters_operators_names=='D_DSM'][0]

    return (adaptated_hyperparameters_operators)
    

def as_text(value):
    if value is None:
        return ""
    return str(value)

def set_column_width(ws,offset=0):
    dim_holder = openpyxl.worksheet.dimensions.DimensionHolder(worksheet=ws)
    j=0

    for col in ws.columns:
        j+=1
        length = max(len(as_text(cell.value)) for cell in col)
        dim_holder[openpyxl.utils.get_column_letter(j)] = openpyxl.worksheet.dimensions.ColumnDimension(ws, min=j, max=j, width=length+offset)

    ws.column_dimensions = dim_holder

def post_traitement(solution,datetime_data,evaluation_function,cost_base,D_movable_load,Y_movable_load,storage_techs,specs_Id,Contract_Id,n_days,file_name_out,Contexte):
                    
    production_set=solution.production_set 
    outputs_solution = evaluation_function(solution,datetime_data,Contexte.storage_characteristics,Contexte.time_resolution,Contexte.n_store,Contexte.duration_years,Contexte.specs_num,Contexte.prices_num,Contexte.fixed_premium,Contexte.Overrun,Contexte.Selling_price,Contexte.Non_movable_load,Contexte.total_D_Movable_load,D_movable_load, Contexte.total_Y_Movable_load,Y_movable_load,Contexte.Grid_Fossil_fuel_ratio,Contexte.Main_grid_PoF_ratio, Contexte.Main_grid_emissions,Contexte.prod_C,Contexte.prods_U,Contexte.Bounds_prod ,Contexte.constraint_num,Contexte.constraint_level,Contexte.cost_constraint,Contexte.n_bits,Contexte.Connexion,Contexte.DG_fuel_consumption,Contexte.DG_fuel_cost,Contexte.DG_unit_cost,Contexte.DG_lifetime,Contexte.DG_maintenance_cost,Contexte.DG_EROI,Contexte.fuel_CO2eq_emissions,storage_techs,n_days)       
    datetime_excel = pd.to_datetime(datetime_data,unit='s').dt.tz_localize(None)

    load = outputs_solution['TimeSeries']['Optimized load (kW)']
    output_baseline = cost_base(solution,datetime_data,Contexte.storage_characteristics,Contexte.time_resolution,Contexte.n_store,Contexte.duration_years,Contexte.specs_num,Contexte.prices_num[0],Contexte.fixed_premium[0],Contexte.Overrun[0],Contexte.Selling_price[0],Contexte.Non_movable_load,D_movable_load, Y_movable_load,Contexte.Grid_Fossil_fuel_ratio,Contexte.Main_grid_PoF_ratio,Contexte.Main_grid_emissions,Contexte.prod_C,Contexte.prods_U,Contexte.Bounds_prod ,Contexte.constraint_num,Contexte.constraint_level,Contexte.cost_constraint,Contexte.n_bits,Contexte.Connexion,Contexte.DG_fuel_consumption,Contexte.DG_fuel_cost,Contexte.DG_unit_cost,Contexte.DG_lifetime,Contexte.DG_maintenance_cost,Contexte.DG_EROI,Contexte.fuel_CO2eq_emissions,storage_techs,n_days)

    
    losses = pd.DataFrame(outputs_solution['TimeSeries']['Losses (kW)'] ,index=[storage_techs[i]+' losses' for i in range(Contexte.n_store)],columns=None).T   
    power_storages = pd.DataFrame(outputs_solution['TimeSeries']['Storage_TS (kW)'],index=[storage_techs[i]+' power (kW)' for i in range(Contexte.n_store)],columns=None).transpose()
    power_storage = power_storages.sum(axis=1)
      
    SOCs=pd.DataFrame(outputs_solution['TimeSeries']['SOCs (%)'],index=[storage_techs[i]+' SOC' for i in range(Contexte.n_store)],columns=None).transpose()   
    production = outputs_solution['TimeSeries']['production (kW)']
    Grid_trading = outputs_solution['TimeSeries']['Grid trading (kW)']

    #Technical
    
    #Environmental
    
    DG_production_solution = outputs_solution['TimeSeries']['DG production (kW)']
        
    curtailment_solution = outputs_solution['TimeSeries']['Curtailment (kW)']

    outputs_TS = pd.DataFrame(data={'Datetime':datetime_excel,'Load (kW)':load,'Power production (kW)':production},index=None)
    outputs_TS = outputs_TS.join(power_storages)
    outputs_TS = outputs_TS.join(losses)
    outputs_TS = outputs_TS.join(pd.DataFrame(data={'Grid power (kW)':Grid_trading,'Grid price (€/kWh)':Contexte.prices_num[solution.contract],'Diesel production (kW)':DG_production_solution,'Curtailment (kW)':curtailment_solution,'Imbalance (kW)':production + power_storage + Grid_trading - load - curtailment_solution + DG_production_solution}))    
    outputs_TS = outputs_TS.join(SOCs)
    

    Technical_outputs=pd.concat((pd.DataFrame(outputs_solution['Technical'],index=['Optimization']),pd.DataFrame(output_baseline['Technical'],index=['Baseline'])))
    Environmental_outputs=pd.concat((pd.DataFrame(outputs_solution['Environment'],index=['Optimization']),pd.DataFrame(output_baseline['Environment'],index=['Baseline'])))
    NPV = outputs_solution['economics']['Value (€)']-output_baseline['economics']['Annual net benefits (€/yrs.)']*outputs_solution['Technical']['Installation lifetime (yrs.)']
    Payback = outputs_solution['economics']['Initial investment (€)']/(outputs_solution['economics']['Annual net benefits (€/yrs.)']-output_baseline['economics']['Annual net benefits (€/yrs.)'])
    if Payback<0 : 
        Payback = np.nan
        
    outputs_solution['economics']["NPV (€)"] = NPV
    outputs_solution['economics']["Payback (yrs.)"] = Payback

    output_baseline['economics']["NPV (€)"] = 0
    output_baseline['economics']["Payback (yrs.)"] = np.nan
    Economic_outputs=pd.concat((pd.DataFrame(outputs_solution['economics'],index=['Optimization']),pd.DataFrame(output_baseline['economics'],index=['Baseline'])))
    
    
    Output_production = pd.DataFrame(data={'ID':specs_Id,'Number of units':production_set,'Coverage ratio':production_set/Contexte.specs_num[:,3],'Initial investment (€)':np.multiply(Contexte.specs_num[:,0],production_set)})
    dist_DOD = outputs_solution['Extra_outputs']['distribution_DOD']
    output_storages = pd.DataFrame(outputs_solution['Storages'],index=storage_techs)
    
    output_useprod = pd.concat((outputs_solution['Extra_outputs']['Uses']['useprod'],output_baseline['Extra_outputs']['Uses']['useprod']))
    output_loadmeet = pd.concat((outputs_solution['Extra_outputs']['Uses']['Loadmeet'],output_baseline['Extra_outputs']['Uses']['Loadmeet']))
    output_whenprod = pd.concat((outputs_solution['Extra_outputs']['Uses']['when_prod'],output_baseline['Extra_outputs']['Uses']['when_prod']))
    output_whenload = pd.concat((outputs_solution['Extra_outputs']['Uses']['when_load'],output_baseline['Extra_outputs']['Uses']['when_load']))
    output_gridexport = pd.concat((outputs_solution['Extra_outputs']['Grid usage']['export'],output_baseline['Extra_outputs']['Grid usage']['export']))
    output_gridimport = pd.concat((outputs_solution['Extra_outputs']['Grid usage']['import'],output_baseline['Extra_outputs']['Grid usage']['import']))

    output_useprod.index,output_loadmeet.index,output_whenprod.index,output_whenload.index,output_gridexport.index,output_gridimport.index = [['Optimization','Baseline'] for i in range(6)]

    
    Flows = pd.concat((pd.DataFrame(outputs_solution['Flows'],index=['Optimization']),pd.DataFrame(output_baseline['Flows'],index=['Baseline'])))
    for i in range(len(storage_techs)):
        Flows = Flows.join(pd.concat((pd.DataFrame(data={storage_techs[i] + ' annual stored energy (kWh)' : outputs_solution['Flows storages']['Annual stored energy (kWh)'][i]},index=['Optimization']),pd.DataFrame(data={storage_techs[i] + ' annual stored energy (kWh)' : output_baseline['Flows storages']['Annual stored energy (kWh)'][i]},index=['Baseline'] ))))
        Flows = Flows.join(pd.concat((pd.DataFrame(data={storage_techs[i] + ' annual reported energy (kWh)' : outputs_solution['Flows storages']['Annual reported energy (kWh)'][i]},index=['Optimization']),pd.DataFrame(data={storage_techs[i] + ' annual reported energy (kWh)' : output_baseline['Flows storages']['Annual reported energy (kWh)'][i]},index=['Baseline'] ))))
        Flows = Flows.join(pd.concat((pd.DataFrame(data={storage_techs[i] + ' annual losses (kWh)' : outputs_solution['Flows storages']['Annual losses (kWh)'][i]},index=['Optimization']),pd.DataFrame(data={storage_techs[i] + ' annual losses (kWh)' : output_baseline['Flows storages']['Annual losses (kWh)'][i]},index=['Baseline'] ))))
    
    if (Contexte.type_optim=='pro'):
        PMS_D_DSM_min_levels = pd.DataFrame({'Hour': np.arange(24)+1, 'min. level (%)': 100*np.concatenate((outputs_solution['PMS']['D_DSM min. levels'],np.array([1.])))})
        PMS_Y_DSM_min_levels = pd.DataFrame({'Month': np.arange(12)+1, 'min. level (%)': 100*np.concatenate((outputs_solution['PMS']['Y_DSM min. levels'],np.array([1.])))})
        discharge_order = pd.DataFrame({'Order':np.arange(Contexte.n_store)+1,'Storage':storage_techs[outputs_solution['PMS']['discharge order']]})
        taking_over = pd.DataFrame({'effective SOC (%)':10*(np.arange(9,0,-1)),'Taking over level (%)':100*outputs_solution['PMS']['taking over'][0],'DG taking over level (%)' if Contexte.Connexion=='Off-grid' else 'Grid taking over level (%)' :100*outputs_solution['PMS']['taking over'][1]})
        
    DG=pd.concat((pd.DataFrame(outputs_solution['DG'],index=['Optimization']),pd.DataFrame(output_baseline['DG'],index=['Baseline'])))
    outputs_solution['Demand-side management']['Load strategy']['Datetime']=datetime_excel
    
    with pd.ExcelWriter(file_name_out,engine='openpyxl') as writer:
        Flows.to_excel(writer,sheet_name='Flows')
        output_useprod.to_excel(writer,sheet_name='Balancing')
        output_loadmeet.to_excel(writer,sheet_name='Balancing',startrow=3)
        output_whenprod.to_excel(writer,sheet_name='Balancing',startrow=6)
        output_whenload.to_excel(writer,sheet_name='Balancing',startrow=9)
        output_gridexport.to_excel(writer,sheet_name='Balancing',startrow=12)
        output_gridimport.to_excel(writer,sheet_name='Balancing',startrow=15)
        Economic_outputs.to_excel(writer,sheet_name='Financial outputs')
        Technical_outputs.to_excel(writer,sheet_name='Technical')        
        pd.DataFrame(output_storages).to_excel(writer,sheet_name='Storages')
        Environmental_outputs.to_excel(writer,sheet_name='Environment outputs')
        dist_DOD.to_excel(writer,sheet_name='SOC distributions',index=None)
        outputs_TS.to_excel(writer,sheet_name='TimeSeries',index=None)
        Output_production.to_excel(writer,sheet_name='Production',index=None)
        if (Contexte.type_optim=='pro'):
            pd.DataFrame(data={'Strategy':outputs_solution['PMS']['strategy']},index=[0]).to_excel(writer,sheet_name='EMS',index=None)
            pd.DataFrame(data={'DSM coefficient':outputs_solution['PMS']['surplus repartition coefficient']},index=[0]).to_excel(writer,sheet_name='EMS',index=None,startrow=3)
            PMS_D_DSM_min_levels.to_excel(writer,sheet_name='EMS',startcol = 2,index=None)
            PMS_Y_DSM_min_levels.to_excel(writer,sheet_name='EMS',startcol = 5,index=None)
            discharge_order.to_excel(writer,sheet_name='EMS',startcol = 8,index=None)
            taking_over.to_excel(writer,sheet_name='EMS',startcol = 11,index=None)
        DG.to_excel(writer,sheet_name='DG')
        outputs_solution['Demand-side management']['Load strategy'].to_excel(writer,sheet_name='DSM',index=None)
        outputs_solution['Demand-side management']['DSM daily strategy'].to_excel(writer,sheet_name='DSM',index=None,startcol=7)
        outputs_solution['Demand-side management']['DSM yearly strategy'].to_excel(writer,sheet_name='DSM',index=None,startcol=13)
        outputs_solution['Balancing']['daily time balancing'].to_excel(writer,sheet_name='Time_balancing',index=None)
        output_baseline['Balancing']['daily time balancing'].to_excel(writer,sheet_name='Time_balancing',index=None,startcol=7)
        outputs_solution['Balancing']['yearly time balancing'].to_excel(writer,sheet_name='Time_balancing',index=None,startcol=14)
        output_baseline['Balancing']['yearly time balancing'].to_excel(writer,sheet_name='Time_balancing',index=None,startcol=21)
        
    #Output Charts
    wb = openpyxl.load_workbook(file_name_out)
    
    set_column_width(wb['Flows'])
    
    ws = wb['Balancing']
    c01 = openpyxl.chart.BarChart()
    c01.title = "Use of production"
    c01.grouping='stacked'
    c01.overlap=100
    c01.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=1, max_col=4, max_row=3), titles_from_data=True)
    c01.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=2, max_col=1, max_row=3))
    c01.x_axis.title = 'Scenario'
    c01.y_axis.title = 'Energy (kWh)'
    c01.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.6, h=0.75))
    c01.layout.layoutTarget = "inner" 
    c01.y_axis.scaling.min = 0
    c01.y_axis.scaling.max = 1.3*max(sum(production),sum(load))/Contexte.time_resolution
    c01.x_axis.delete = False
    c01.y_axis.delete = False
    c01.height = 7.5
    c01.width = 20
    c01.legend.overlay = False
    c01.x_axis.tickLblPos = "low"
    c01.x_axis.tickLblPos = "low"
    ws.add_chart(c01, "A1")   
    
    c02 = openpyxl.chart.BarChart()
    c02.title = "Load response"
    c02.grouping='stacked'
    c02.overlap=100
    c02.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=4, max_col=4, max_row=6), titles_from_data=True)
    c02.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=5, max_col=1, max_row=6))
    c02.x_axis.title = 'Scenario'
    c02.y_axis.title = 'Energy (kWh)'
    c02.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.6, h=0.75))
    c02.layout.layoutTarget = "inner" 
    c02.y_axis.scaling.min = 0
    c02.y_axis.scaling.max = 1.3*max(sum(production),sum(load))/Contexte.time_resolution
    c02.x_axis.delete = False
    c02.y_axis.delete = False
    c02.height = 7.5
    c02.width = 20
    c02.legend.overlay = False
    c02.x_axis.tickLblPos = "low"
    c02.x_axis.tickLblPos = "low"
    ws.add_chart(c02, "F1")   
    
    c03 = openpyxl.chart.BarChart()
    c03.title = "When production exceeds load"
    c03.grouping='stacked'
    c03.overlap=100
    c03.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=7, max_col=4+2*Contexte.n_store, max_row=9), titles_from_data=True)
    c03.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=8, max_col=1, max_row=9))
    c03.x_axis.title = 'Scenario'
    c03.y_axis.title = 'Energy (kWh)'
    c03.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.6, h=0.75))
    c03.layout.layoutTarget = "inner" 
    c03.y_axis.scaling.min = 0
    c03.x_axis.delete = False
    c03.y_axis.delete = False
    c03.height = 7.5
    c03.width = 20
    c03.legend.overlay = False
    c03.x_axis.tickLblPos = "low"
    c03.x_axis.tickLblPos = "low"
    ws.add_chart(c03, "A31") 
    
    c04 = openpyxl.chart.BarChart()
    c04.title = "When load exceeds production"
    c04.grouping='stacked'
    c04.overlap=100
    c04.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=10, max_col=4+2*Contexte.n_store, max_row=12), titles_from_data=True)
    c04.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=11, max_col=1, max_row=12))
    c04.x_axis.title = 'Scenario'
    c04.y_axis.title = 'Energy (kWh)'
    c04.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.6, h=0.75))
    c04.layout.layoutTarget = "inner" 
    c04.y_axis.scaling.min = 0
    c04.x_axis.delete = False
    c04.y_axis.delete = False
    c04.height = 7.5
    c04.width = 20
    c04.legend.overlay = False
    c04.x_axis.tickLblPos = "low"
    c04.x_axis.tickLblPos = "low"
    ws.add_chart(c04, "F31") 
    
    c05 = openpyxl.chart.BarChart()
    c05.title = "Grid usage (exportation)"
    c05.grouping='stacked'
    c05.overlap=100
    c05.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=13, max_col=3, max_row=15), titles_from_data=True)
    c05.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=14, max_col=1, max_row=15))
    c05.x_axis.title = 'Scenario'
    c05.y_axis.title = 'Energy (kWh)'
    c05.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.6, h=0.75))
    c05.layout.layoutTarget = "inner" 
    c05.y_axis.scaling.min = 0
    c05.x_axis.delete = False
    c05.y_axis.delete = False
    c05.height = 7.5
    c05.width = 20
    c05.legend.overlay = False
    c05.x_axis.tickLblPos = "low"
    c05.x_axis.tickLblPos = "low"
    ws.add_chart(c05, "A16") 
    
    c06 = openpyxl.chart.BarChart()
    c06.title = "Grid usage (importation)"
    c06.grouping='stacked'
    c06.overlap=100
    c06.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=16, max_col=3, max_row=18), titles_from_data=True)
    c06.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=17, max_col=1, max_row=18))
    c06.x_axis.title = 'Scenario'
    c06.y_axis.title = 'Energy (kWh)'
    c06.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.6, h=0.75))
    c06.layout.layoutTarget = "inner" 
    c06.y_axis.scaling.min = 0
    c06.x_axis.delete = False
    c06.y_axis.delete = False
    c06.height = 7.5
    c06.width = 20
    c06.legend.overlay = False
    c06.x_axis.tickLblPos = "low"
    c06.x_axis.tickLblPos = "low"
    ws.add_chart(c06, "F16") 
    
    set_column_width(ws)
    
    ws = wb['Financial outputs']
    c2 = openpyxl.chart.BarChart()
    c2.title = "LCOE decomposition"
    c2.grouping='stacked'
    c2.overlap=100
    c2.add_data(openpyxl.chart.Reference(ws,min_col=3, min_row=1, max_col=12, max_row=3), titles_from_data=True)
    c2.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=2, max_col=1, max_row=3))
    c2.x_axis.title = 'Scenario'
    c2.y_axis.title = 'Cost (€/kWh)'
    c2.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c2.layout.layoutTarget = "inner"  
    c2.x_axis.delete = False
    c2.y_axis.delete = False
    c2.legend.position = 'b'  
    c2.height = 15
    c2.width = 15
    c2.legend.overlay = False
    c2.x_axis.tickLblPos = "low"
    c2.x_axis.tickLblPos = "low"
    ws.add_chart(c2, "B6")  
    
    c2_2 = openpyxl.chart.BarChart()
    c2_2.title = "Exportation income"
    c2_2.grouping='stacked'
    c2_2.overlap=100
    c2_2.add_data(openpyxl.chart.Reference(ws,min_col=13, min_row=1, max_col=13, max_row=3), titles_from_data=True)
    c2_2.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=2, max_col=1, max_row=3))
    c2_2.x_axis.title = 'Scenario'
    c2_2.y_axis.title = 'Income (€/kWh)'
    c2_2.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c2_2.layout.layoutTarget = "inner"  
    c2_2.x_axis.delete = False
    c2_2.y_axis.delete = False
    c2_2.legend.position = 'b'
    c2_2.height = 15
    c2_2.width = 15
    c2_2.legend.overlay = False
    c2_2.x_axis.tickLblPos = "low"
    c2_2.x_axis.tickLblPos = "low"
    ws.add_chart(c2_2, "F6")   
    
    set_column_width(ws)
    
    set_column_width(wb['Technical'])
    set_column_width(wb['Environment outputs'])  
    set_column_width(wb['Storages'])
    
    ws = wb['SOC distributions']
    c1 = openpyxl.chart.LineChart()
    c1.title = "SOC percentile distributions"
    c1.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=1, max_col=1+Contexte.n_store, max_row=101), titles_from_data=True)
    c1.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=2, max_col=1, max_row=102))
    c1.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c1.layout.layoutTarget = "inner"    
    c1.height,c1.width = (10,18)
    c1.x_axis.scaling.min,c1.y_axis.scaling.min = (0,0)
    c1.x_axis.delete = False
    c1.y_axis.delete = False
    c1.legend.position = 'b'
    c1.x_axis.title = 'Percentile'
    c1.y_axis.title = 'Density'
    ws.add_chart(c1, "F2") 
    
    set_column_width(ws)
    #fmt = '%Y-%m-%d %H:%M:%S'
   
    ws = wb['TimeSeries']
    c3 = openpyxl.chart.ScatterChart()
    c3.title = "Storages SOCs"
    xvalues = openpyxl.chart.Reference(ws,min_col=1, min_row=1, max_col=1, max_row=Contexte.n_bits+1)
    for i in range(Contexte.n_store):
        values = openpyxl.chart.Reference(ws,min_col=9+Contexte.n_store*2+i, min_row=1, max_col=9+Contexte.n_store*2+i, max_row=Contexte.n_bits+1)
        series = openpyxl.chart.Series(values, xvalues, title_from_data=True)
        series.graphicalProperties.line.width = 12000
        c3.series.append(series)
    c3.y_axis.title = 'SOC'
    c3.x_axis.title = 'Date'
    c3.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.0015, y=0.005,w=0.97, h=0.7))
    c3.layout.layoutTarget = "inner"    
    c3.x_axis.delete = True
    c3.y_axis.delete = False
    c3.legend.position = 'l'
    (c3.height,c3.width) = (9,200)
    c3.legend.overlay = False
    c3.x_axis.majorUnit = Contexte.time_resolution*24*7
    c3.y_axis.scaling.min,c3.y_axis.scaling.max = (0.0,1.0)    
    c3.y_axis.majorUnit = 0.2
    c3.x_axis.scaling.min,c3.x_axis.scaling.max =(0,Contexte.n_bits)
    ws.add_chart(c3, "O3") 
    
    set_column_width(ws)
    
    c4 = openpyxl.chart.ScatterChart()
    c4.title = "TimeSeries"
    xvalues = openpyxl.chart.Reference(ws,min_col=1, min_row=1, max_col=1, max_row=Contexte.n_bits+1)
    values1 = openpyxl.chart.Reference(ws,min_col=2, min_row=1, max_col=2, max_row=Contexte.n_bits+1)
    series1 = openpyxl.chart.Series(values1, xvalues, title_from_data=True)
    series1.graphicalProperties.line.width = 12000
    values2 = openpyxl.chart.Reference(ws,min_col=3, min_row=1, max_col=3, max_row=Contexte.n_bits+1)
    series2 = openpyxl.chart.Series(values2, xvalues, title_from_data=True)
    series2.graphicalProperties.line.width = 12000
    values3 = openpyxl.chart.Reference(ws,min_col=6+2*Contexte.n_store, min_row=1, max_col=6+2*Contexte.n_store, max_row=Contexte.n_bits+1)
    series3 = openpyxl.chart.Series(values3, xvalues, title_from_data=True)
    series3.graphicalProperties.line.width = 12000
    c4.series.append(series1)
    c4.series.append(series2)
    c4.series.append(series3)    
    c4.x_axis.title = 'Date'
    c4.y_axis.title = 'Power (kW)'
    c4.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.0015, y=0.005,w=0.97, h=0.7))
    c4.layout.layoutTarget = "inner"    
    c4.x_axis.delete = True
    c4.y_axis.delete = False
    c4.legend.position = 'l'
    (c4.height,c4.width) = (9,200)
    c4.legend.overlay = False
    c4.x_axis.majorUnit = Contexte.time_resolution*24*7
    c4.x_axis.scaling.min,c4.x_axis.scaling.max =(0,Contexte.n_bits)
    ws.add_chart(c4, "O23")  
    
    c5 = openpyxl.chart.ScatterChart()
    c5.title = "Storage timeSeries"
    xvalues = openpyxl.chart.Reference(ws,min_col=1, min_row=1, max_col=1, max_row=Contexte.n_bits+1)
    for i in range(Contexte.n_store) :
        values = openpyxl.chart.Reference(ws,min_col=4+i, min_row=1, max_col=4+i, max_row=Contexte.n_bits+1)
        series = openpyxl.chart.Series(values, xvalues, title_from_data=True)
        series.graphicalProperties.line.width = 12000
        c5.series.append(series) 
    c5.x_axis.title = 'Date'
    c5.y_axis.title = 'Power (kW)'
    c5.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.0015, y=0.005,w=0.97, h=0.7))
    c5.layout.layoutTarget = "inner"    
    c5.x_axis.delete = True
    c5.y_axis.delete = False
    c5.legend.position = 'l'
    (c5.height,c5.width) = (9,200)
    c5.legend.overlay = False
    c5.x_axis.majorUnit = Contexte.time_resolution*24*7
    c5.x_axis.scaling.min,c5.x_axis.scaling.max =(0,Contexte.n_bits)
    ws.add_chart(c5, "O43")  
    
    set_column_width(ws)
    
    ws = wb['Production']
    cProd = openpyxl.chart.BarChart()
    cProd.title = "Production set ratios"
    cProd.overlap=30
    cProd.add_data(openpyxl.chart.Reference(ws,min_col=3, min_row=1, max_col=3, max_row=len(production_set)+1), titles_from_data=True)
    cProd.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=2, max_col=1, max_row=len(production_set)+1))
    cProd.x_axis.title = 'Production ID'
    cProd.y_axis.title = 'Coverage ratio'
    cProd.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.01, y=0.005,w=0.85, h=0.7))
    cProd.layout.layoutTarget = "inner"  
    cProd.x_axis.delete = False
    cProd.y_axis.delete = False
    cProd.legend = None
    cProd.height = 8
    cProd.width = 26
    #cProd.legend.overlay = False
    cProd.x_axis.tickLblPos = "low"
    cProd.x_axis.tickLblPos = "low"
    cProd.series[0].graphicalProperties.solidFill = "ff9900" 
    cProd.series[0].graphicalProperties.line.solidFill = "00000" 
    cProd.y_axis.scaling.min,cProd.y_axis.scaling.max =(0,1)
    ws.add_chart(cProd, "F2")   
    
    set_column_width(ws)
    
    ws = wb['DSM']
    c6 = openpyxl.chart.LineChart()
    c6.title = "Daily Demand-side management"
    c6.add_data(openpyxl.chart.Reference(ws,min_col=9, min_row=1, max_col=12, max_row=int(Contexte.time_resolution*24)+1), titles_from_data=True)
    c6.set_categories(openpyxl.chart.Reference(ws,min_col=8, min_row=2, max_col=8, max_row=int(Contexte.time_resolution*24)+1))
    c6.x_axis.title = 'Datetime'
    c6.y_axis.title = 'Load (kW)'
    c6.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c6.layout.layoutTarget = "inner"    
    c6.height,c6.width = (10,18)
    c6.x_axis.delete = False
    c6.y_axis.delete = False
    c6.legend.position = 'b'
    c6.legend.overlay = False
    c6.tickLblPos = "low"
    c6.x_axis.tickLblPos = "low"
    ws.add_chart(c6, "D10") 
    
    c7 = openpyxl.chart.LineChart()
    c7.title = "Yearly Demand-side management"
    c7.add_data(openpyxl.chart.Reference(ws,min_col=15, min_row=1, max_col=18, max_row=int(n_days)+1), titles_from_data=True)
    c7.set_categories(openpyxl.chart.Reference(ws,min_col=14, min_row=2, max_col=14, max_row=int(n_days)+1))
    c7.x_axis.title = 'Datetime'
    c7.y_axis.title = 'Load (kW)'
    c7.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c7.layout.layoutTarget = "inner"    
    c7.height,c7.width = (10,18)
    c7.x_axis.delete = False
    c7.y_axis.delete = False
    c7.legend.position = 'b'
    c7.legend.overlay = False
    c7.tickLblPos = "low"
    c7.x_axis.tickLblPos = "low"
    ws.add_chart(c7, "N10") 
    
    set_column_width(ws)
    
    ws = wb['Time_balancing']
    c8 = openpyxl.chart.LineChart()
    c8.title = "Daily balancing (optimization)"
    c8.add_data(openpyxl.chart.Reference(ws,min_col=2, min_row=1, max_col=6, max_row=int(Contexte.time_resolution*24)+1), titles_from_data=True)
    c8.set_categories(openpyxl.chart.Reference(ws,min_col=1, min_row=2, max_col=1, max_row=int(Contexte.time_resolution*24)))
    c8.x_axis.title = 'Datetime'
    c8.y_axis.title = 'Power (kW)'
    c8.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c8.layout.layoutTarget = "inner"    
    c8.height,c8.width = (10,18)
    c8.x_axis.delete = False
    c8.y_axis.delete = False
    c8.legend.position = 'b'
    c8.legend.overlay = False
    c8.tickLblPos = "low"
    c8.x_axis.tickLblPos = "low"
    ws.add_chart(c8, "A10") 
    
    c9 = openpyxl.chart.LineChart()
    c9.title = "Daily balancing (baseline)"
    c9.add_data(openpyxl.chart.Reference(ws,min_col=9, min_row=1, max_col=13, max_row=int(Contexte.time_resolution*24)+1), titles_from_data=True)
    c9.set_categories(openpyxl.chart.Reference(ws,min_col=8, min_row=2, max_col=8, max_row=int(Contexte.time_resolution*24)))
    c9.x_axis.title = 'Datetime'
    c9.y_axis.title = 'Power (kW)'
    c9.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c9.layout.layoutTarget = "inner"    
    c9.height,c9.width = (10,18)
    c9.x_axis.delete = False
    c9.y_axis.delete = False
    c9.legend.position = 'b'
    c9.legend.overlay = False
    c9.tickLblPos = "low"
    c9.x_axis.tickLblPos = "low"
    ws.add_chart(c9, "H10") 
    
    c10 = openpyxl.chart.LineChart()
    c10.title = "Yearly balancing (optimization)"
    c10.add_data(openpyxl.chart.Reference(ws,min_col=16, min_row=1, max_col=20, max_row=int(n_days)+1), titles_from_data=True)
    c10.set_categories(openpyxl.chart.Reference(ws,min_col=15, min_row=2, max_col=15, max_row=int(n_days)))
    c10.x_axis.title = 'Datetime'
    c10.y_axis.title = 'Power (kW)'
    c10.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c10.layout.layoutTarget = "inner"    
    c10.height,c10.width = (10,18)
    c10.x_axis.delete = False
    c10.y_axis.delete = False
    c10.legend.position = 'b'
    c10.legend.overlay = False
    c10.tickLblPos = "low"
    c10.x_axis.tickLblPos = "low"
    ws.add_chart(c10, "O10") 
    
    c11 = openpyxl.chart.LineChart()
    c11.title = "Yearly balancing (baseline)"
    c11.add_data(openpyxl.chart.Reference(ws,min_col=23, min_row=1, max_col=27, max_row=int(n_days)+1), titles_from_data=True)
    c11.set_categories(openpyxl.chart.Reference(ws,min_col=22, min_row=2, max_col=22, max_row=int(n_days)))
    c11.x_axis.title = 'Datetime'
    c11.y_axis.title = 'Power (kW)'
    c11.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
    c11.layout.layoutTarget = "inner"    
    c11.height,c11.width = (10,18)
    c11.x_axis.delete = False
    c11.y_axis.delete = False
    c11.legend.position = 'b'
    c11.legend.overlay = False
    c11.tickLblPos = "low"
    c11.x_axis.tickLblPos = "low"
    ws.add_chart(c11, "V10") 
    
    set_column_width(ws)
    
    set_column_width(wb['DG'],offset=2)
    
    if  (Contexte.type_optim=='pro'):
        ws = wb['EMS']
    
        c12_1 = openpyxl.chart.LineChart()
        c12_1.title = "Taking over"
        data = openpyxl.chart.Reference(ws, min_col=13, min_row=1, max_col=14, max_row=10)
        c12_1.add_data(data, titles_from_data=True)
        c12_1.set_categories(openpyxl.chart.Reference(ws,min_col=12, min_row=2, max_col=12, max_row=10))
        for i in range(2):
            c12_1.series[i].marker.symbol = "square"
            c12_1.series[i].smooth = False
        c12_1.x_axis.title = 'effective SOC (%)'
        c12_1.y_axis.title = 'Secondary storage involvment (%)'
        c12_1.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
        c12_1.layout.layoutTarget = "inner"
        c12_1.x_axis.delete = False
        c12_1.y_axis.delete = False
        c12_1.legend.position = 'b'
        c12_1.legend.overlay = False
        c12_1.tickLblPos = "low"
        c12_1.x_axis.tickLblPos = "low"
        c12_1.x_axis.scaling.min,c12_1.y_axis.scaling.min,c12_1.y_axis.scaling.max = (0,0,100)
        ws.add_chart(c12_1, "G20") 
    
        c12_2 = openpyxl.chart.ScatterChart()
        c12_2.title = "D_DSM minimum levels"    
        xvalues=openpyxl.chart.Reference(ws,min_col=3, min_row=2, max_col=3, max_row=25)
        values = openpyxl.chart.Reference(ws, min_col=4, min_row=2, max_col=4, max_row=25)
        series = openpyxl.chart.Series(values, xvalues, title='Minimum level (%)')
        series.marker=openpyxl.chart.marker.Marker('square')
        series.smooth = False
        c12_2.varyColors = False
        c12_2.series.append(series)  
        c12_2.x_axis.title = 'Hour'
        c12_2.y_axis.title = 'minimum level (%)'
        c12_2.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
        c12_2.layout.layoutTarget = "inner"
        c12_2.x_axis.delete = False
        c12_2.y_axis.delete = False
        c12_2.legend.position = 'b'
        c12_2.legend.overlay = False
        c12_2.tickLblPos = "low"
        c12_2.x_axis.tickLblPos = "low"
        c12_2.x_axis.scaling.min,c12_2.x_axis.scaling.max,c12_2.y_axis.scaling.min,c12_2.y_axis.scaling.max = (1,24,0,100)

        ws.add_chart(c12_2, "P3") 
    
    
        c12_3 = openpyxl.chart.ScatterChart()
        c12_3.title = "Y_DSM minimum levels"
        xvalues=openpyxl.chart.Reference(ws,min_col=6, min_row=2, max_col=6, max_row=13)
        values = openpyxl.chart.Reference(ws, min_col=7, min_row=2, max_col=7, max_row=13)
        series = openpyxl.chart.Series(values, xvalues, title='Minimum level (%)')
        series.marker=openpyxl.chart.marker.Marker('square')
        series.smooth = False
        c12_3.varyColors = False
        c12_3.series.append(series)  
        c12_3.x_axis.title = 'Month'
        c12_3.y_axis.title = 'minimum level (%)'
        c12_3.layout = openpyxl.chart.layout.Layout(manualLayout=openpyxl.chart.layout.ManualLayout(x=0.025, y=0.005,w=0.85, h=0.7))
        c12_3.layout.layoutTarget = "inner"
        c12_3.x_axis.delete = False
        c12_3.y_axis.delete = False
        c12_3.legend.position = 'b'
        c12_3.legend.overlay = False
        c12_3.tickLblPos = "low"
        c12_3.x_axis.tickLblPos = "low"
        c12_3.x_axis.scaling.min,c12_3.x_axis.scaling.max,c12_3.y_axis.scaling.min,c12_3.y_axis.scaling.max = (1,12,0,100)
        ws.add_chart(c12_3, "P20") 
    
        set_column_width(wb['EMS'],offset=2)

    
    
    wb.save(file_name_out)
  
    